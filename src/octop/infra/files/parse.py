"""Bounded structured extraction for conversation attachments.

The knowledge-base parser keeps its existing flat-string contract. This module
adds source locations for chat reads without changing its callers.
"""

from __future__ import annotations

import csv
import io
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from octop.infra.knowledge.parse import _parse_html, _parse_json, _read_text

PARSER_VERSION = 2
_OFFICE_EXTENSIONS = frozenset({".docx", ".xlsx", ".xlsm", ".pptx"})


class FileParseError(ValueError):
    """The attachment cannot be safely or completely interpreted."""


class UnsupportedFileTypeError(FileParseError):
    """No parser is available for the attachment format."""


@dataclass(frozen=True)
class ParseLimits:
    max_input_bytes: int = 20 * 1024 * 1024
    max_pages: int = 200
    max_cells: int = 100_000
    max_characters: int = 200_000
    max_segments: int = 2_000
    max_archive_entries: int = 10_000
    max_archive_uncompressed_bytes: int = 100 * 1024 * 1024
    max_archive_ratio: int = 100


DEFAULT_PARSE_LIMITS = ParseLimits()


@dataclass(frozen=True)
class FileSegment:
    text: str
    locator: dict[str, Any]


@dataclass(frozen=True)
class ParseCoverage:
    complete: bool
    characters: int
    total_units: int | None
    processed_units: int
    empty_units: int
    truncated: bool
    parser_version: int = PARSER_VERSION
    ocr_needed_units: int = 0


@dataclass(frozen=True)
class FileParseResult:
    segments: tuple[FileSegment, ...]
    coverage: ParseCoverage


class _Collector:
    def __init__(self, limits: ParseLimits) -> None:
        self.limits = limits
        self.segments: list[FileSegment] = []
        self.characters = 0
        self.processed_units = 0
        self.empty_units = 0
        self.ocr_needed_units = 0
        self.truncated = False

    def add(self, text: str, locator: dict[str, Any]) -> bool:
        cleaned = text.strip()
        if not cleaned:
            self.empty_units += 1
            return True
        if len(self.segments) >= self.limits.max_segments:
            self.truncated = True
            return False
        remaining = self.limits.max_characters - self.characters
        if len(cleaned) > remaining:
            self.truncated = True
            return False
        self.segments.append(FileSegment(cleaned, locator))
        self.characters += len(cleaned)
        return True

    def result(self, *, total_units: int | None) -> FileParseResult:
        return FileParseResult(
            segments=tuple(self.segments),
            coverage=ParseCoverage(
                complete=not self.truncated and self.ocr_needed_units == 0,
                characters=self.characters,
                total_units=total_units,
                processed_units=self.processed_units,
                empty_units=self.empty_units,
                truncated=self.truncated,
                ocr_needed_units=self.ocr_needed_units,
            ),
        )


def _validate_archive(path: Path, limits: ParseLimits) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > limits.max_archive_entries:
                raise FileParseError("office archive entry limit exceeded")
            total = 0
            seen: set[str] = set()
            for entry in entries:
                name = entry.filename.replace("\\", "/")
                parts = name.rstrip("/").split("/")
                if (
                    name.startswith("/")
                    or "//" in name
                    or any(part in {"", ".", ".."} for part in parts)
                ):
                    raise FileParseError("office archive contains an invalid path")
                if name in seen:
                    raise FileParseError("office archive contains duplicate entries")
                seen.add(name)
                if entry.flag_bits & 0x1:
                    raise FileParseError("encrypted office archives are unsupported")
                total += entry.file_size
                if total > limits.max_archive_uncompressed_bytes:
                    raise FileParseError("office archive expanded size limit exceeded")
                if entry.file_size and (
                    entry.compress_size == 0
                    or entry.file_size > entry.compress_size * limits.max_archive_ratio
                ):
                    raise FileParseError("office archive compression ratio limit exceeded")
    except FileParseError:
        raise
    except (OSError, zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise FileParseError("invalid office archive") from exc


def _plain_text(path: Path, collector: _Collector) -> FileParseResult:
    text = _read_text(path)
    lines = text.splitlines()
    for line_number, line in enumerate(lines, start=1):
        if not collector.add(line, {"kind": "text_line", "start": line_number, "end": line_number}):
            break
        collector.processed_units += 1
    if len(lines) > collector.processed_units:
        collector.truncated = True
    return collector.result(total_units=len(lines))


def _delimited(path: Path, collector: _Collector, delimiter: str) -> FileParseResult:
    rows = csv.reader(io.StringIO(_read_text(path), newline=""), delimiter=delimiter)
    total_units: int | None = 0
    cells = 0
    for row_number, row in enumerate(rows, start=1):
        cells += len(row)
        if cells > collector.limits.max_cells:
            collector.truncated = True
            total_units = None
            break
        text = "\t".join(cell.replace("\n", " ").replace("\r", " ").strip() for cell in row).strip()
        if not collector.add(text, {"kind": "csv_row", "row": row_number}):
            total_units = None
            break
        collector.processed_units += 1
    if total_units is not None:
        total_units = collector.processed_units
    return collector.result(total_units=total_units)


def _pdf(path: Path, collector: _Collector) -> FileParseResult:
    from pypdf import PdfReader

    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise FileParseError("encrypted PDFs are unsupported")
        total = len(reader.pages)
        for page_index, page in enumerate(reader.pages):
            if page_index >= collector.limits.max_pages:
                collector.truncated = True
                break
            text = page.extract_text() or ""
            if not text.strip() and page.images:
                collector.ocr_needed_units += 1
            if not collector.add(text, {"kind": "pdf_page", "page": page_index + 1}):
                break
            collector.processed_units += 1
        if total > collector.limits.max_pages:
            collector.truncated = True
        return collector.result(total_units=total)
    except FileParseError:
        raise
    except Exception as exc:
        raise FileParseError("invalid PDF document") from exc


def _xml_text(element: Any) -> str:
    from docx.oxml.ns import qn

    parts: list[str] = []
    for child in element.iter():
        if child.tag == qn("w:t"):
            parts.append(child.text or "")
        elif child.tag == qn("w:tab"):
            parts.append("\t")
        elif child.tag == qn("w:br"):
            parts.append("\n")
    return "".join(parts)


def _docx(path: Path, collector: _Collector) -> FileParseResult:
    from docx import Document
    from docx.oxml.ns import qn

    try:
        document = Document(str(path))
        for block_number, block in enumerate(document.element.body.iterchildren(), start=1):
            if block.tag == qn("w:p"):
                paragraphs = [block]
            elif block.tag == qn("w:tbl"):
                paragraphs = []
            else:
                paragraphs = block.xpath(".//w:p")
            if block.tag == qn("w:tbl"):
                for row_index, row in enumerate(block.xpath("./w:tr"), start=1):
                    if not collector.add(
                        "\t".join(_xml_text(cell) for cell in row.xpath("./w:tc")),
                        {"kind": "docx_table_row", "table": block_number, "row": row_index},
                    ):
                        break
                    collector.processed_units += 1
            else:
                for paragraph in paragraphs:
                    if not collector.add(
                        _xml_text(paragraph),
                        {"kind": "docx_paragraph", "block": block_number},
                    ):
                        break
                    collector.processed_units += 1
            if collector.truncated:
                break
        return collector.result(total_units=collector.processed_units)
    except FileParseError:
        raise
    except Exception as exc:
        raise FileParseError("invalid Word document") from exc


def _xlsx(path: Path, collector: _Collector) -> FileParseResult:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            total_units: int | None = 0
            total_cells = 0
            for sheet in workbook.worksheets:
                max_row, max_column = sheet.max_row, sheet.max_column
                if max_row is None or max_column is None:
                    collector.truncated = True
                    total_units = None
                    break
                sheet_cells = max_row * max_column
                if total_cells + sheet_cells > collector.limits.max_cells:
                    collector.truncated = True
                    total_units = None
                    break
                total_cells += sheet_cells
                total_units += max_row
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    for column_number, value in enumerate(row, start=1):
                        if value is None:
                            continue
                        text = str(value).replace("\t", " ").replace("\n", " ").strip()
                        if not text:
                            continue
                        if not collector.add(
                            text,
                            {
                                "kind": "spreadsheet_cell",
                                "sheet": sheet.title,
                                "cell": f"{get_column_letter(column_number)}{row_number}",
                            },
                        ):
                            total_units = None
                            break
                    collector.processed_units += 1
                    if collector.truncated:
                        break
                if collector.truncated:
                    break
            return collector.result(total_units=total_units)
        finally:
            workbook.close()
    except FileParseError:
        raise
    except Exception as exc:
        raise FileParseError("invalid spreadsheet") from exc


def _pptx(path: Path, collector: _Collector) -> FileParseResult:
    from pptx import Presentation

    try:
        presentation = Presentation(str(path))
        total = len(presentation.slides)
        for slide_index, slide in enumerate(presentation.slides):
            if slide_index >= collector.limits.max_pages:
                collector.truncated = True
                break
            for shape_index, shape in enumerate(slide.shapes, start=1):
                if hasattr(shape, "text") and not collector.add(
                    shape.text,
                    {"kind": "pptx_shape", "slide": slide_index + 1, "shape": shape_index},
                ):
                    break
            collector.processed_units += 1
            if collector.truncated:
                break
        return collector.result(total_units=total)
    except FileParseError:
        raise
    except Exception as exc:
        raise FileParseError("invalid presentation") from exc


def parse_file(path: Path, *, limits: ParseLimits = DEFAULT_PARSE_LIMITS) -> FileParseResult:
    """Extract bounded text segments with stable format-specific source locators."""
    path = Path(path)
    if not path.is_file():
        raise FileParseError("file does not exist")
    try:
        if path.stat().st_size > limits.max_input_bytes:
            raise FileParseError("file size limit exceeded")
    except OSError as exc:
        raise FileParseError("file cannot be read") from exc

    suffix = path.suffix.lower()
    if suffix in _OFFICE_EXTENSIONS:
        _validate_archive(path, limits)
    collector = _Collector(limits)

    if suffix in {
        ".txt",
        ".md",
        ".markdown",
        ".rst",
        ".json",
        ".jsonl",
        ".yaml",
        ".yml",
        ".toml",
        ".ini",
        ".py",
        ".ts",
        ".tsx",
        ".js",
        ".jsx",
        ".html",
        ".htm",
        ".xml",
        ".sql",
        ".sh",
        ".log",
        ".diff",
        ".patch",
    }:
        if suffix in {".html", ".htm"}:
            text = _parse_html(path)
        elif suffix == ".json":
            text = _parse_json(path)
        else:
            text = _read_text(path)
        # Use the shared decoder while keeping the extracted lines as locators.
        decoded = text.splitlines()
        for line_number, line in enumerate(decoded, start=1):
            if not collector.add(
                line, {"kind": "text_line", "start": line_number, "end": line_number}
            ):
                break
            collector.processed_units += 1
        if len(decoded) > collector.processed_units:
            collector.truncated = True
        return collector.result(total_units=len(decoded))
    if suffix == ".csv":
        return _delimited(path, collector, ",")
    if suffix == ".tsv":
        return _delimited(path, collector, "\t")
    if suffix == ".pdf":
        return _pdf(path, collector)
    if suffix == ".docx":
        return _docx(path, collector)
    if suffix in {".xlsx", ".xlsm"}:
        return _xlsx(path, collector)
    if suffix == ".pptx":
        return _pptx(path, collector)
    raise UnsupportedFileTypeError(f"unsupported document format: {suffix or '(none)'}")
