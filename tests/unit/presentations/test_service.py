from __future__ import annotations

import asyncio
import io
import json
from contextlib import asynccontextmanager
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from pptx import Presentation

from octop.infra.errors import ErrorCode, OctopError
from octop.infra.presentations.models import CreateDeck, JobRequest, SaveDeck, Slide
from octop.infra.presentations.service import PresentationService
from octop.infra.presentations.svg import validate_svg

SVG = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"><rect width="1280" height="720" fill="#ffffff"/><text x="80" y="100" font-size="40">Hello</text></svg>'


class MemorySettings:
    def __init__(self) -> None:
        self.values: dict[str, str] = {}

    def get(self, key: str) -> str | None:
        return self.values.get(key)

    def set(self, key: str, value: str) -> None:
        self.values[key] = value

    def list_prefix(self, prefix: str) -> dict[str, str]:
        return {key: value for key, value in self.values.items() if key.startswith(prefix)}

    def compare_and_set(self, key: str, expected: str, value: str) -> bool:
        if self.values.get(key) != expected:
            return False
        self.values[key] = value
        return True


@pytest.fixture
def service() -> PresentationService:
    return PresentationService(
        SimpleNamespace(
            services=SimpleNamespace(
                settings_repo=MemorySettings(),
                user_policy_repo=SimpleNamespace(get=Mock(return_value=None)),
                usage_repo=SimpleNamespace(
                    record=Mock(return_value=1), total_tokens_for_user=Mock(return_value=0)
                ),
            )
        )
    )


@pytest.mark.asyncio
async def test_persists_scoped_decks_and_rejects_stale_save(service: PresentationService) -> None:
    deck = await service.create_deck(7, "agent-a", CreateDeck(title="Plan", brief="Our plan"))
    assert [item.deck_id for item in await service.list_decks(7, "agent-a")] == [deck.deck_id]
    assert await service.list_decks(8, "agent-a") == []
    with pytest.raises(OctopError) as denied:
        await service.get_deck(8, "agent-a", deck.deck_id)
    assert denied.value.status == 404
    assert denied.value.to_envelope(locale="en")["error"]["message"] == "Presentation not found."
    body = SaveDeck(
        revision=deck.revision,
        title="Changed",
        brief=deck.brief,
        audience=deck.audience,
        style=deck.style,
        slides=[],
    )
    updated = await service.save_deck(7, "agent-a", deck.deck_id, body)
    assert updated.title == "Changed" and updated.revision > deck.revision
    with pytest.raises(OctopError) as stale:
        await service.save_deck(7, "agent-a", deck.deck_id, body)
    assert stale.value.status == 409


@pytest.mark.asyncio
async def test_save_only_invalidates_changed_page(service: PresentationService) -> None:
    deck = await service.create_deck(7, "agent-a", CreateDeck(title="Plan"))
    slides = [
        Slide(slide_id="first", title="First", body="Text", svg=SVG),
        Slide(slide_id="second", title="Second", body="Text", svg=SVG),
    ]
    deck.slides = slides
    await service._persist(deck)
    summaries = await service.list_decks(7, "agent-a")
    assert summaries[0].slides == [] and summaries[0].page_count == deck.page_count
    changed = [s.model_copy(deep=True) for s in slides]
    changed[0].title = "New title"
    # Clients cannot supply replacement SVG, even for a page they own.
    changed[0].svg = '<svg onload="alert(1)"/>'
    updated = await service.save_deck(
        7,
        "agent-a",
        deck.deck_id,
        SaveDeck(
            revision=deck.revision,
            title=deck.title,
            brief=deck.brief,
            audience=deck.audience,
            style=deck.style,
            slides=changed,
        ),
    )
    assert updated.slides[0].svg is None
    assert updated.slides[1].svg == SVG


@pytest.mark.parametrize(
    "payload",
    [
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"><script>alert(1)</script></svg>',
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720" onload="alert(1)"/>',
        '<!DOCTYPE svg [<!ENTITY x "abc">]><svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"/>',
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"><image href="https://example.com/a.png"/></svg>',
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720"><rect fill="url(https://example.com/a)"/></svg>',
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 999999999 720"/>',
    ],
)
def test_rejects_active_or_unbounded_svg(payload: str) -> None:
    with pytest.raises(ValueError):
        validate_svg(payload)


def test_accepts_safe_svg_and_bounds_size() -> None:
    assert validate_svg(SVG) == SVG
    with pytest.raises(ValueError):
        validate_svg(SVG + " " * 110_000)


@pytest.mark.parametrize("replacement", ['width="-1"', 'width="1e300"', 'width="NaN"'])
def test_rejects_invalid_geometry(replacement: str) -> None:
    with pytest.raises(ValueError):
        validate_svg(SVG.replace('width="1280"', replacement))


def test_rejects_svg_depth_and_unknown_namespace() -> None:
    with pytest.raises(ValueError):
        validate_svg(
            '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1280 720">'
            + "<g>" * 40
            + "</g>" * 40
            + "</svg>"
        )
    with pytest.raises(ValueError):
        validate_svg(
            SVG.replace("<text ", '<foreign:text xmlns:foreign="https://example.com" ').replace(
                "</text>", "</foreign:text>"
            )
        )


@pytest.mark.asyncio
async def test_failure_preserves_completed_pages_and_recovers_after_restart(
    service: PresentationService, monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = await service.create_deck(7, "agent-a", CreateDeck(title="Plan"))
    deck.slides = [
        Slide(slide_id="first", title="First", body="Text", svg=SVG),
        Slide(slide_id="second", title="Second", body="Text"),
    ]
    await service._persist(deck)
    monkeypatch.setattr(
        service, "_invoke", AsyncMock(side_effect=RuntimeError("private provider error"))
    )
    busy = await service.start_job(
        7, "agent-a", deck.deck_id, JobRequest(revision=deck.revision, action="generate")
    )
    assert busy.status == "generating"
    with pytest.raises(OctopError) as duplicate:
        await service.start_job(
            7, "agent-a", deck.deck_id, JobRequest(revision=busy.revision, action="generate")
        )
    assert duplicate.value.status == 409
    for _ in range(100):
        await asyncio.sleep(0.01)
        failed = await service.get_deck(7, "agent-a", deck.deck_id)
        if failed.status == "failed":
            break
    assert failed.status == "failed"
    assert failed.slides[0].svg == SVG
    assert "private provider" not in (failed.error or "")
    failed.status = "generating"
    await service._persist(failed)
    restarted = PresentationService(service.server)
    recovered = await restarted.get_deck(7, "agent-a", deck.deck_id)
    assert recovered.status == "failed" and recovered.slides[0].svg == SVG
    await service.close()


@pytest.mark.asyncio
async def test_generate_checkpoints_then_revision_is_single_page(
    service: PresentationService, monkeypatch: pytest.MonkeyPatch
) -> None:
    deck = await service.create_deck(7, "agent-a", CreateDeck(title="Plan"))
    deck.slides = [
        Slide(slide_id="first", title="First", body="Text"),
        Slide(slide_id="second", title="Second", body="Text"),
    ]
    await service._persist(deck)
    invoke = AsyncMock(return_value=SVG)
    monkeypatch.setattr(service, "_invoke", invoke)
    await service.start_job(
        7, "agent-a", deck.deck_id, JobRequest(revision=deck.revision, action="generate")
    )
    for _ in range(100):
        await asyncio.sleep(0.01)
        deck = await service.get_deck(7, "agent-a", deck.deck_id)
        if deck.status == "idle":
            break
    assert all(s.svg == SVG for s in deck.slides) and invoke.await_count == 2
    invoke.reset_mock()
    revised_svg = SVG.replace("Hello", "Revised")
    invoke.return_value = json.dumps(
        {"title": "Updated", "body": "Revised text", "notes": "Speaker notes", "svg": revised_svg}
    )
    await service.start_job(
        7,
        "agent-a",
        deck.deck_id,
        JobRequest(
            revision=deck.revision, action="revise", page_id="second", instruction="Change title"
        ),
    )
    for _ in range(100):
        await asyncio.sleep(0.01)
        deck = await service.get_deck(7, "agent-a", deck.deck_id)
        if deck.status == "idle":
            break
    assert deck.slides[0].svg == SVG and deck.slides[1].svg == revised_svg
    assert deck.slides[1].title == "Updated" and deck.slides[1].notes == "Speaker notes"
    assert invoke.await_count == 1
    await service.close()


@pytest.mark.asyncio
async def test_generate_retries_svg_rejected_by_the_platform_validator(
    service: PresentationService,
) -> None:
    deck = await service.create_deck(7, "agent-a", CreateDeck(title="Plan"))
    deck.slides = [Slide(slide_id="first", title="First", body="Text")]
    await service._persist(deck)
    invalid_svg = SVG.replace("<rect", '<rect style="font-family:Arial"')
    invoke = AsyncMock(side_effect=[invalid_svg, SVG])
    service._invoke = invoke  # type: ignore[method-assign]

    await service.start_job(
        7, "agent-a", deck.deck_id, JobRequest(revision=deck.revision, action="generate")
    )
    for _ in range(100):
        await asyncio.sleep(0.01)
        deck = await service.get_deck(7, "agent-a", deck.deck_id)
        if deck.status in {"idle", "failed"}:
            break

    assert deck.status == "idle"
    assert deck.slides[0].svg == SVG
    assert invoke.await_count == 2
    assert "failed platform validation" in invoke.await_args_list[1].args[2]
    assert "unsupported SVG attribute" in invoke.await_args_list[1].args[2]
    await service.close()


@pytest.mark.asyncio
async def test_generate_survives_reasoning_only_responses_until_third_attempt(
    service: PresentationService,
) -> None:
    deck = await service.create_deck(7, "agent-a", CreateDeck(title="Plan"))
    deck.slides = [Slide(slide_id="first", title="First", body="Text")]
    await service._persist(deck)
    reasoning = "1. Analyze the request.\n2. Draft: <!DOCTYPE html> sketch without artwork."
    invoke = AsyncMock(side_effect=[reasoning, reasoning, SVG])
    service._invoke = invoke  # type: ignore[method-assign]

    await service.start_job(
        7, "agent-a", deck.deck_id, JobRequest(revision=deck.revision, action="generate")
    )
    for _ in range(100):
        await asyncio.sleep(0.01)
        deck = await service.get_deck(7, "agent-a", deck.deck_id)
        if deck.status in {"idle", "failed"}:
            break

    assert deck.status == "idle"
    assert deck.slides[0].svg == SVG
    assert invoke.await_count == 3
    assert "no SVG found in response" in invoke.await_args_list[1].args[2]
    assert "no SVG found in response" in invoke.await_args_list[2].args[2]
    await service.close()


@pytest.mark.asyncio
async def test_compose_retries_an_invalid_outline_and_reports_outline_failure(
    service: PresentationService,
) -> None:
    deck = await service.create_deck(7, "agent-a", CreateDeck(title="Plan", page_count=1))
    invoke = AsyncMock(return_value='{"slides":[]}')
    service._invoke = invoke  # type: ignore[method-assign]

    await service.start_job(
        7, "agent-a", deck.deck_id, JobRequest(revision=deck.revision, action="compose")
    )
    for _ in range(100):
        await asyncio.sleep(0.01)
        deck = await service.get_deck(7, "agent-a", deck.deck_id)
        if deck.status in {"idle", "failed"}:
            break

    assert deck.status == "failed"
    assert deck.error == "The model did not return a valid presentation outline."
    assert invoke.await_count == 3
    await service.close()


@pytest.mark.asyncio
async def test_compose_retries_outline_then_renders_pages(service: PresentationService) -> None:
    deck = await service.create_deck(7, "agent-a", CreateDeck(title="Plan", page_count=1))
    outline = '{"slides":[{"title":"First","body":"Text","notes":""}]}'
    invoke = AsyncMock(side_effect=['{"slides":[]}', outline, SVG])
    service._invoke = invoke  # type: ignore[method-assign]

    await service.start_job(
        7, "agent-a", deck.deck_id, JobRequest(revision=deck.revision, action="compose")
    )
    for _ in range(100):
        await asyncio.sleep(0.01)
        deck = await service.get_deck(7, "agent-a", deck.deck_id)
        if deck.status in {"idle", "failed"}:
            break

    assert deck.status == "idle"
    assert len(deck.slides) == 1 and deck.slides[0].svg == SVG
    assert invoke.await_count == 3
    await service.close()


@pytest.mark.asyncio
@pytest.mark.parametrize("running", [True, False])
async def test_export_uploads_real_package_and_scopes_download(
    service: PresentationService, monkeypatch: pytest.MonkeyPatch, running: bool
) -> None:
    binary = io.BytesIO()
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(binary)
    files: dict[str, bytes] = {}

    async def upload(path: str, data: bytes) -> None:
        files[path] = data

    workspace = SimpleNamespace(
        aupload_bytes=AsyncMock(side_effect=upload),
        adownload_bytes=AsyncMock(side_effect=lambda path: files.get(path)),
    )
    live_agent = Mock(return_value=SimpleNamespace(workspace=workspace))
    if not running:
        live_agent.side_effect = OctopError(ErrorCode.AGENT_NOT_RUNNING, "stopped")
    service.server.app_runtime = SimpleNamespace(
        agent_registry=SimpleNamespace(
            get_agent=live_agent,
            workspace_for_agent=lambda _id: workspace,
        )
    )
    monkeypatch.setattr(
        "octop.infra.presentations.service.runtime_status", lambda: {"installed": True}
    )
    exporter = Mock(return_value=binary.getvalue())
    monkeypatch.setattr("octop.infra.presentations.service.export_native", exporter)
    deck = await service.create_deck(7, "agent-a", CreateDeck(title="Plan"))
    deck.slides = [
        Slide(slide_id="first", title="First", body="Text", notes="Speaker note", svg=SVG)
    ]
    await service._persist(deck)
    await service.start_job(
        7, "agent-a", deck.deck_id, JobRequest(revision=deck.revision, action="export")
    )
    for _ in range(100):
        await asyncio.sleep(0.01)
        deck = await service.get_deck(7, "agent-a", deck.deck_id)
        if deck.status == "idle":
            break
    assert deck.export_path and deck.export_path.startswith(
        f"presentations/7/{deck.deck_id}/exports/"
    )
    assert exporter.call_args.args[0][0].notes == "Speaker note"
    assert await service.download(7, "agent-a", deck.deck_id) == binary.getvalue()
    with pytest.raises(OctopError):
        await service.download(8, "agent-a", deck.deck_id)
    await service.close()


@pytest.mark.asyncio
async def test_bound_model_accepts_text_content_blocks(service: PresentationService) -> None:
    @asynccontextmanager
    async def tracked(_agent_id: str):
        yield

    async def _chunk_agen():
        yield SimpleNamespace(content=[{"type": "text", "text": SVG}])

    model = Mock()
    model.bind.return_value = model
    model.astream = Mock(return_value=_chunk_agen())
    registry = SimpleNamespace(
        get_row=lambda _id: SimpleNamespace(default_model="provider/model"),
        get_config=lambda _id: {},
        providers=SimpleNamespace(
            is_model_ref_usable=lambda _ref: True, get_model_reasoning_capability=lambda _ref: None
        ),
        harness_manager=SimpleNamespace(
            shared_factory=SimpleNamespace(get_chat_model=lambda _ref: model)
        ),
        _track_invocation=tracked,
    )
    service.server.app_runtime = SimpleNamespace(agent_registry=registry)
    deck = await service.create_deck(7, "agent-a", CreateDeck(title="Plan"))
    assert (
        await service._invoke(
            "agent-a",
            7,
            "System instructions",
            {"title": "Plan"},
            thread_id=f"presentation:{deck.deck_id}",
        )
        == SVG
    )
    assert model.astream.call_count == 1


def _bound_runtime(service: PresentationService, response: object) -> Mock:
    @asynccontextmanager
    async def tracked(_agent_id: str):
        yield

    async def _chunk_agen():
        yield response

    model = Mock()
    model.bind.return_value = model
    model.astream = Mock(return_value=_chunk_agen())
    registry = SimpleNamespace(
        get_row=lambda _id: SimpleNamespace(default_model="provider/model"),
        get_config=lambda _id: {},
        providers=SimpleNamespace(
            is_model_ref_usable=lambda _ref: True, get_model_reasoning_capability=lambda _ref: None
        ),
        harness_manager=SimpleNamespace(
            shared_factory=SimpleNamespace(get_chat_model=lambda _ref: model)
        ),
        _track_invocation=tracked,
    )
    service.server.app_runtime = SimpleNamespace(agent_registry=registry)
    return model


@pytest.mark.asyncio
async def test_zero_token_quota_prevents_model_call(service: PresentationService) -> None:
    model = _bound_runtime(service, SimpleNamespace(content=SVG))
    service.server.services.user_policy_repo.get.return_value = "0"
    await service.create_deck(7, "agent-a", CreateDeck(title="Plan"))
    with pytest.raises(OctopError) as denied:
        await service._invoke("agent-a", 7, "Instructions", {}, thread_id="presentation:new")
    assert denied.value.code is ErrorCode.TOKEN_QUOTA_EXCEEDED
    model.astream.assert_not_called()
    service.server.services.usage_repo.record.assert_not_called()


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "metadata",
    [
        {
            "usage_metadata": {
                "input_tokens": 120,
                "output_tokens": 30,
                "input_token_details": {"cache_read": 20},
                "output_token_details": {"reasoning": 5},
            }
        },
        {
            "response_metadata": {
                "token_usage": {
                    "prompt_tokens": 120,
                    "completion_tokens": 30,
                    "prompt_tokens_details": {"cached_tokens": 20},
                    "completion_tokens_details": {"reasoning_tokens": 5},
                }
            }
        },
        {
            "usage": {
                "prompt_tokens": 120,
                "completion_tokens": 30,
                "prompt_tokens_details": {"cached_tokens": 20},
                "completion_tokens_details": {"reasoning_tokens": 5},
            }
        },
    ],
)
async def test_successful_model_call_records_real_usage(
    service: PresentationService, metadata: dict[str, object]
) -> None:
    _bound_runtime(service, SimpleNamespace(content=SVG, **metadata))
    deck = await service.create_deck(7, "agent-a", CreateDeck(title="Plan"))
    await service._invoke(
        "agent-a", 7, "Instructions", {}, thread_id=f"presentation:{deck.deck_id}"
    )
    record = service.server.services.usage_repo.record
    record.assert_called_once()
    values = record.call_args.kwargs
    assert values["agent_id"] == "agent-a" and values["user_id"] == 7
    assert values["model"] == "provider/model"
    assert values["source"] == "presentation_studio"
    assert values["thread_id"] == f"presentation:{deck.deck_id}"
    assert values["input_tokens"] == 120 and values["output_tokens"] == 30
    assert values["cache_read_tokens"] == 20 and values["uncached_input_tokens"] == 100
    assert values["reasoning_tokens"] == 5 and values["model_calls"] == 1
