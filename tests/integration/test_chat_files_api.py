"""Integration tests for the conversation file read surface (T04)."""

from __future__ import annotations

import asyncio
import io
import zipfile
from typing import Any

from docx import Document

from tests.support.auth import create_user


async def _upload_and_bind(
    client: Any,
    srv: Any,
    auth: dict[str, str],
    agent_id: str,
    *,
    filename: str = "note.txt",
    data: bytes = b"alpha beta\n\ngamma beta\n",
    thread_id: str = "files-thread",
    read_file: bool = True,
) -> dict[str, Any]:
    upload = await client.post(
        f"/api/agents/{agent_id}/upload",
        files={"file": (filename, data, "text/plain")},
        headers=auth,
    )
    assert upload.status_code == 200, upload.text
    body = upload.json()
    record = srv.services.chat_file_repo.get(body["file_id"])
    assert record is not None
    srv.services.thread_repo.insert(
        thread_id=thread_id,
        agent_id=agent_id,
        user_id=record.user_id,
        channel_type="dashboard",
        session_key=thread_id,
    )
    assert srv.services.chat_file_repo.bind_thread(
        record.file_id,
        user_id=record.user_id,
        agent_id=agent_id,
        thread_id=thread_id,
        revision_id=body["revision_id"],
        workspace_path=body["workspace_path"],
    )
    if read_file:
        read = await client.post(
            f"/api/agents/{agent_id}/files/{record.file_id}/read",
            headers=auth,
            json={"thread_id": thread_id, "revision_id": body["revision_id"]},
        )
        assert read.status_code == 200, read.text
    return {**body, "file_id": record.file_id, "thread_id": thread_id}


async def test_upload_reports_file_identity_size_and_read_status(
    env_with_agent: Any,
) -> None:
    client, srv, auth, aid = env_with_agent
    body = await _upload_and_bind(client, srv, auth, aid)
    assert body["file_id"]
    assert body["revision_id"]
    assert int(body["size"]) == len(b"alpha beta\n\ngamma beta\n")
    assert body["read_status"] == "stored"


async def test_text_artifact_revision_uses_cas_and_keeps_stale_revision(
    env_with_agent: Any,
) -> None:
    client, srv, auth, aid = env_with_agent
    agent = srv.services.agent_repo.get(aid)
    assert agent is not None and agent.user_id is not None
    thread_id = "generated-text-thread"
    srv.services.thread_repo.insert(
        thread_id=thread_id,
        agent_id=aid,
        user_id=agent.user_id,
        channel_type="dashboard",
        session_key=thread_id,
    )

    created = await client.post(
        f"/api/agents/{aid}/threads/{thread_id}/files",
        headers=auth,
        json={"filename": "draft.md", "content": "# First\n"},
    )
    assert created.status_code == 201, created.text
    first = created.json()
    assert first["display_name"] == "draft.md"
    assert first["revisions"][0]["origin_kind"] == "generated"
    assert first["access_url"]
    first_download = await client.get(first["access_url"], headers=auth)
    assert first_download.status_code == 200
    assert first_download.content == b"# First\n"
    text_url = (
        f"/api/agents/{aid}/files/{first['file_id']}/revisions/{first['current_revision_id']}/text"
    )
    text_revision = await client.get(
        text_url,
        params={"thread_id": thread_id},
        headers=auth,
    )
    assert text_revision.status_code == 200
    assert text_revision.json()["content"] == "# First\n"
    wrong_thread = await client.get(
        text_url,
        params={"thread_id": "another-thread"},
        headers=auth,
    )
    assert wrong_thread.status_code == 404
    intruder = await create_user(client, auth, username="artifact-download-intruder")
    srv.services.agent_repo.set_shared(aid, True)
    denied_download = await client.get(first["access_url"], headers=intruder)
    assert denied_download.status_code == 404, denied_download.text
    first_revision = first["current_revision_id"]

    url = f"/api/agents/{aid}/files/{first['file_id']}/revisions"
    attempts = await asyncio.gather(
        client.post(
            url,
            headers=auth,
            json={
                "thread_id": thread_id,
                "base_revision_id": first_revision,
                "content": "# Updated\n",
            },
        ),
        client.post(
            url,
            headers=auth,
            json={
                "thread_id": thread_id,
                "base_revision_id": first_revision,
                "content": "# Competing edit\n",
            },
        ),
    )
    updated = next(result for result in attempts if result.status_code == 201)
    stale = next(result for result in attempts if result.status_code == 409)
    latest_revision = updated.json()["current_revision_id"]
    assert latest_revision != first_revision
    assert stale.json()["error"]["code"] == "FILE_REVISION_CONFLICT"
    conflict_revision = stale.json()["error"]["details"]["revision_id"]
    current = await client.get(f"/api/agents/{aid}/files/{first['file_id']}", headers=auth)
    assert current.status_code == 200
    assert current.json()["current_revision_id"] == latest_revision
    assert len(current.json()["revisions"]) == 3
    conflict_item = next(
        revision
        for revision in current.json()["revisions"]
        if revision["revision_id"] == conflict_revision
    )
    assert conflict_item["access_url"]

    read = await client.post(
        f"/api/agents/{aid}/files/{first['file_id']}/read",
        headers=auth,
        json={"thread_id": thread_id, "revision_id": latest_revision},
    )
    assert read.status_code == 200, read.text
    winning_text = "# Updated" if attempts.index(updated) == 0 else "# Competing edit"
    assert winning_text in "".join(segment["text"] for segment in read.json()["segments"])
    branch_read = await client.post(
        f"/api/agents/{aid}/files/{first['file_id']}/read",
        headers=auth,
        json={"thread_id": thread_id, "revision_id": conflict_revision},
    )
    assert branch_read.status_code == 200, branch_read.text
    losing_text = "# Competing edit" if attempts.index(updated) == 0 else "# Updated"
    assert losing_text in "".join(segment["text"] for segment in branch_read.json()["segments"])
    branch_download = await client.get(conflict_item["access_url"], headers=auth)
    assert branch_download.status_code == 200
    assert branch_download.content.decode("utf-8").strip() == losing_text

    latest_row = srv.services.chat_file_repo.get_revision(latest_revision)
    assert latest_row is not None
    live_agent = srv.app_runtime.agent_registry.get_agent(aid)
    await live_agent.workspace.aupload_bytes(latest_row.workspace_path, b"tampered")
    corrupt_download = await client.get(current.json()["access_url"], headers=auth)
    assert corrupt_download.status_code == 409
    assert corrupt_download.json()["error"]["code"] == "FILE_INTEGRITY_FAILED"


async def test_generated_docx_is_editable_as_markdown_and_uploaded_docx_is_not(
    env_with_agent: Any,
) -> None:
    client, srv, auth, aid = env_with_agent
    agent = srv.services.agent_repo.get(aid)
    assert agent is not None and agent.user_id is not None
    thread_id = "generated-docx-thread"
    srv.services.thread_repo.insert(
        thread_id=thread_id,
        agent_id=aid,
        user_id=agent.user_id,
        channel_type="dashboard",
        session_key=thread_id,
    )

    created = await client.post(
        f"/api/agents/{aid}/threads/{thread_id}/files",
        headers=auth,
        json={"filename": "report.docx", "content": "# Report\n\nInitial text"},
    )
    assert created.status_code == 201, created.text
    file = created.json()
    assert (
        file["media_type"]
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    thread_files = await client.get(f"/api/agents/{aid}/threads/{thread_id}/files", headers=auth)
    assert thread_files.status_code == 200, thread_files.text
    assert thread_files.json()[0]["origin_kind"] == "generated"
    download = await client.get(file["access_url"], headers=auth)
    assert download.status_code == 200
    assert [p.text for p in Document(io.BytesIO(download.content)).paragraphs] == [
        "Report",
        "Initial text",
    ]
    too_large = await client.post(
        f"/api/agents/{aid}/threads/{thread_id}/files",
        headers=auth,
        json={"filename": "large.docx", "content": "x" * 200_001},
    )
    assert too_large.json()["error"]["code"] == "DOCUMENT_CONVERSION_LIMIT"
    assert too_large.json()["error"]["details"]["max_characters"] == 200_000
    too_many_lines = await client.post(
        f"/api/agents/{aid}/threads/{thread_id}/files",
        headers=auth,
        json={"filename": "long.docx", "content": "a\n" * 10_001},
    )
    assert too_many_lines.json()["error"]["code"] == "DOCUMENT_CONVERSION_LIMIT"
    assert too_many_lines.json()["error"]["details"]["max_lines"] == 10_000
    invalid_xml = await client.post(
        f"/api/agents/{aid}/threads/{thread_id}/files",
        headers=auth,
        json={"filename": "invalid.docx", "content": "bad\x01xml"},
    )
    assert invalid_xml.json()["error"]["code"] == "ATTACHMENT_UNSUPPORTED_TYPE"

    read_url = (
        f"/api/agents/{aid}/files/{file['file_id']}/revisions/{file['current_revision_id']}/text"
    )
    readable = await client.get(read_url, params={"thread_id": thread_id}, headers=auth)
    assert readable.status_code == 200, readable.text
    assert "# Report" in readable.json()["content"]
    assert "Initial text" in readable.json()["content"]

    revised = await client.post(
        f"/api/agents/{aid}/files/{file['file_id']}/revisions",
        headers=auth,
        json={
            "thread_id": thread_id,
            "base_revision_id": file["current_revision_id"],
            "content": "# Report\n\nRevised text",
        },
    )
    assert revised.status_code == 201, revised.text
    latest = revised.json()
    latest_download = await client.get(latest["access_url"], headers=auth)
    assert latest_download.status_code == 200
    assert [p.text for p in Document(io.BytesIO(latest_download.content)).paragraphs] == [
        "Report",
        "Revised text",
    ]
    preserved_original = await client.get(file["access_url"], headers=auth)
    assert preserved_original.status_code == 200
    assert [p.text for p in Document(io.BytesIO(preserved_original.content)).paragraphs] == [
        "Report",
        "Initial text",
    ]

    uploaded = await _upload_and_bind(
        client,
        srv,
        auth,
        aid,
        filename="uploaded.docx",
        data=download.content,
        thread_id="uploaded-docx-thread",
        read_file=False,
    )
    uploaded_list = await client.get(
        f"/api/agents/{aid}/threads/{uploaded['thread_id']}/files", headers=auth
    )
    assert uploaded_list.status_code == 200, uploaded_list.text
    assert uploaded_list.json()[0]["origin_kind"] == "upload"
    uploaded_text = await client.get(
        f"/api/agents/{aid}/files/{uploaded['file_id']}/revisions/{uploaded['revision_id']}/text",
        params={"thread_id": uploaded["thread_id"]},
        headers=auth,
    )
    assert uploaded_text.json()["error"]["code"] == "ATTACHMENT_UNSUPPORTED_TYPE"


async def test_unsupported_file_type_is_reported_and_persisted(env_with_agent: Any) -> None:
    client, srv, auth, aid = env_with_agent
    body = await _upload_and_bind(
        client,
        srv,
        auth,
        aid,
        filename="archive.bin",
        data=b"not a supported document",
        read_file=False,
    )

    response = await client.post(
        f"/api/agents/{aid}/files/{body['file_id']}/read",
        headers=auth,
        json={"thread_id": body["thread_id"], "revision_id": body["revision_id"]},
    )

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "ATTACHMENT_UNSUPPORTED_TYPE"
    assert (
        srv.services.chat_file_repo.get_revision(body["revision_id"]).read_status == "unsupported"
    )


async def test_invalid_document_is_reported_as_read_failure(env_with_agent: Any) -> None:
    client, srv, auth, aid = env_with_agent
    body = await _upload_and_bind(
        client,
        srv,
        auth,
        aid,
        filename="broken.pdf",
        data=b"not a pdf",
        read_file=False,
    )

    response = await client.post(
        f"/api/agents/{aid}/files/{body['file_id']}/read",
        headers=auth,
        json={"thread_id": body["thread_id"], "revision_id": body["revision_id"]},
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "FILE_READ_FAILED"
    assert srv.services.chat_file_repo.get_revision(body["revision_id"]).read_status == "failed"


async def test_thread_files_list_and_detail_hide_host_paths(env_with_agent: Any) -> None:
    client, srv, auth, aid = env_with_agent
    body = await _upload_and_bind(client, srv, auth, aid)

    listing = await client.get(f"/api/agents/{aid}/threads/{body['thread_id']}/files", headers=auth)
    assert listing.status_code == 200, listing.text
    assert listing.headers["cache-control"] == "private, no-store"
    rows = listing.json()
    assert [row["file_id"] for row in rows] == [body["file_id"]]
    assert rows[0]["display_name"] == "note.txt"
    assert rows[0]["read_status"] == "ready"
    assert "workspace_path" not in rows[0]

    detail = await client.get(f"/api/agents/{aid}/files/{body['file_id']}", headers=auth)
    assert detail.status_code == 200, detail.text
    payload = detail.json()
    assert payload["current_revision_id"] == body["revision_id"]
    assert payload["coverage"]["complete"] is True
    assert [item["revision_id"] for item in payload["revisions"]] == [body["revision_id"]]
    assert "workspace_path" not in payload
    assert all("workspace_path" not in item for item in payload["revisions"])


async def test_segments_paginate_with_next_offset(env_with_agent: Any) -> None:
    client, srv, auth, aid = env_with_agent
    body = await _upload_and_bind(client, srv, auth, aid)

    first = await client.get(
        f"/api/agents/{aid}/files/{body['file_id']}/segments",
        params={"limit": 1},
        headers=auth,
    )
    assert first.status_code == 200, first.text
    page = first.json()
    assert [item["ordinal"] for item in page["segments"]] == [0]
    assert page["next_offset"] == 1
    assert "workspace_path" not in page["segments"][0]
    span = page["segments"][0]["source_span"]
    revision = srv.services.chat_file_repo.get_revision(body["revision_id"])
    assert span == {
        "file_id": body["file_id"],
        "revision_id": body["revision_id"],
        "parser_version": 2,
        "content_hash": revision.sha256,
        "ordinal": 0,
        "locator": {"kind": "text_line", "start": 1, "end": 1},
    }

    second = await client.get(
        f"/api/agents/{aid}/files/{body['file_id']}/segments",
        params={"limit": 1, "offset": 1},
        headers=auth,
    )
    assert second.status_code == 200, second.text
    assert [item["ordinal"] for item in second.json()["segments"]] == [1]
    assert second.json()["next_offset"] is None


async def test_search_returns_bounded_case_insensitive_matches(env_with_agent: Any) -> None:
    client, srv, auth, aid = env_with_agent
    body = await _upload_and_bind(client, srv, auth, aid)

    found = await client.post(
        f"/api/agents/{aid}/files/{body['file_id']}/search",
        headers=auth,
        json={"query": "BETA"},
    )
    assert found.status_code == 200, found.text
    matches = found.json()["matches"]
    assert [item["ordinal"] for item in matches] == [0, 1]
    assert all("workspace_path" not in item for item in matches)

    wildcard = await client.post(
        f"/api/agents/{aid}/files/{body['file_id']}/search",
        headers=auth,
        json={"query": "%"},
    )
    assert wildcard.status_code == 200, wildcard.text
    assert wildcard.json()["matches"] == []


async def test_read_surface_rejects_unbound_draft(env_with_agent: Any) -> None:
    client, _srv, auth, aid = env_with_agent
    upload = await client.post(
        f"/api/agents/{aid}/upload",
        files={"file": ("draft.txt", b"draft body", "text/plain")},
        headers=auth,
    )
    assert upload.status_code == 200, upload.text
    file_id = upload.json()["file_id"]

    # Metadata is the owner's own data and stays available for drafts.
    detail = await client.get(f"/api/agents/{aid}/files/{file_id}", headers=auth)
    assert detail.status_code == 200, detail.text
    assert detail.json()["current_revision_id"] == upload.json()["revision_id"]

    # Derived body content requires the thread binding, like read does.
    for call in (
        client.get(f"/api/agents/{aid}/files/{file_id}/segments", headers=auth),
        client.post(
            f"/api/agents/{aid}/files/{file_id}/search",
            headers=auth,
            json={"query": "draft"},
        ),
    ):
        response = await call
        assert response.status_code == 404, response.text


async def test_workspace_mutations_cannot_change_registered_file_revision(
    env_with_agent: Any,
) -> None:
    client, srv, auth, aid = env_with_agent
    body = await _upload_and_bind(client, srv, auth, aid)
    revision = srv.services.chat_file_repo.get_revision(body["revision_id"])
    assert revision is not None
    path = f"/{revision.workspace_path}"

    overwritten = await client.put(
        f"/api/agents/{aid}/workspace/file",
        params={"path": path, "from_workspace": "true"},
        headers=auth,
        json={"content": "tampered replacement"},
    )
    assert overwritten.status_code == 403, overwritten.text

    deleted = await client.delete(
        f"/api/agents/{aid}/workspace/file",
        params={"path": path, "from_workspace": "true"},
        headers=auth,
    )
    assert deleted.status_code == 403, deleted.text

    moved = await client.post(
        f"/api/agents/{aid}/workspace/move",
        params={"path": path, "from_workspace": "true"},
        headers=auth,
        json={"destination": "/inbound/renamed.txt"},
    )
    assert moved.status_code == 403, moved.text

    read = await client.get(
        f"/api/agents/{aid}/workspace/download",
        params={"path": path, "from_workspace": "true"},
        headers=auth,
    )
    assert read.status_code == 200, read.text
    assert read.content == b"alpha beta\n\ngamma beta\n"


async def test_legacy_workspace_reads_enforce_chat_file_owner_and_revision_hash(
    env_with_agent: Any,
) -> None:
    client, srv, auth, aid = env_with_agent
    body = await _upload_and_bind(client, srv, auth, aid)
    revision = srv.services.chat_file_repo.get_revision(body["revision_id"])
    assert revision is not None
    path = f"/{revision.workspace_path}"
    own_url = f"/api/agents/{aid}/workspace/download"
    own_download = await client.get(
        own_url,
        params={"path": path, "from_workspace": "true"},
        headers=auth,
    )
    assert own_download.status_code == 200
    assert own_download.content == b"alpha beta\n\ngamma beta\n"

    picture = await client.post(
        f"/api/agents/{aid}/upload",
        files={"file": ("private.png", b"private image bytes", "image/png")},
        headers=auth,
    )
    assert picture.status_code == 200, picture.text
    intruder = await create_user(client, auth, username="legacy-workspace-intruder")
    srv.services.agent_repo.set_shared(aid, True)
    for endpoint in (
        own_url,
        f"/api/agents/{aid}/workspace/file",
    ):
        response = await client.get(
            endpoint,
            params={"path": path, "from_workspace": "true"},
            headers=intruder,
        )
        assert response.status_code == 404, response.text
    preview = await client.get(
        f"/api/agents/{aid}/media/preview",
        params={"source": picture.json()["workspace_path"], "mime_type": "image/png"},
        headers=intruder,
    )
    assert preview.status_code == 404, preview.text

    workspace = srv.app_runtime.agent_registry.get_agent(aid).workspace
    await workspace.aupload_bytes(revision.workspace_path, b"tampered")
    corrupted = await client.get(
        own_url,
        params={"path": path, "from_workspace": "true"},
        headers=auth,
    )
    assert corrupted.status_code == 409, corrupted.text
    assert corrupted.json()["error"]["code"] == "FILE_INTEGRITY_FAILED"


async def test_soft_deleted_file_is_hidden_from_old_workspace_download_url(
    env_with_agent: Any,
) -> None:
    client, srv, auth, aid = env_with_agent
    uploaded = await client.post(
        f"/api/agents/{aid}/upload",
        files={"file": ("old-draft.txt", b"old draft", "text/plain")},
        headers=auth,
    )
    assert uploaded.status_code == 200, uploaded.text
    payload = uploaded.json()
    file_row = srv.services.chat_file_repo.get(payload["file_id"])
    assert file_row is not None
    revision = srv.services.chat_file_repo.get_revision(payload["revision_id"])
    assert revision is not None
    deleted = await client.delete(
        f"/api/agents/{aid}/files/{payload['file_id']}/draft", headers=auth
    )
    assert deleted.status_code == 204
    workspace = srv.app_runtime.agent_registry.workspace_for_agent(aid)
    assert workspace is not None
    assert await workspace.adownload_bytes(revision.workspace_path) is None

    old_download = await client.get(
        f"/api/agents/{aid}/workspace/download",
        params={"path": f"/{revision.workspace_path}", "from_workspace": "true"},
        headers=auth,
    )
    assert old_download.status_code == 404, old_download.text


async def test_workspace_archive_excludes_another_shared_users_chat_file(
    env_with_agent: Any,
) -> None:
    client, srv, auth, aid = env_with_agent
    own = await client.post(
        f"/api/agents/{aid}/upload",
        files={"file": ("owner-file.txt", b"owner content", "text/plain")},
        headers=auth,
    )
    assert own.status_code == 200, own.text
    intruder = await create_user(client, auth, username="workspace-archive-intruder")
    srv.services.agent_repo.set_shared(aid, True)
    foreign = await client.post(
        f"/api/agents/{aid}/upload",
        files={"file": ("private-file.txt", b"foreign content", "text/plain")},
        headers=intruder,
    )
    assert foreign.status_code == 200, foreign.text

    tree = await client.get(
        f"/api/agents/{aid}/workspace/tree",
        params={"path": "/inbound", "from_workspace": "true"},
        headers=auth,
    )
    assert tree.status_code == 200, tree.text
    listed_paths = {row["path"] for row in tree.json()}
    assert any(path.endswith("_owner-file.txt") for path in listed_paths)
    assert all(not path.endswith("_private-file.txt") for path in listed_paths)
    globbed = await client.get(
        f"/api/agents/{aid}/workspace/glob",
        params={"pattern": "*.txt", "path": "/inbound", "from_workspace": "true"},
        headers=auth,
    )
    assert globbed.status_code == 200, globbed.text
    assert all("private-file.txt" not in str(row) for row in globbed.json())
    grep = await client.get(
        f"/api/agents/{aid}/workspace/grep",
        params={"pattern": "foreign content", "path": "/inbound", "from_workspace": "true"},
        headers=auth,
    )
    assert grep.status_code == 200, grep.text
    assert all("foreign content" not in str(row) for row in grep.json())

    archive = await client.get(f"/api/agents/{aid}/workspace/archive", headers=auth)

    assert archive.status_code == 200, archive.text
    with zipfile.ZipFile(io.BytesIO(archive.content)) as bundle:
        names = set(bundle.namelist())
        assert own.json()["workspace_path"] in names
        assert foreign.json()["workspace_path"] not in names
        assert b"foreign content" not in b"".join(bundle.read(name) for name in names)


async def test_workspace_archive_import_preserves_registered_chat_file_revisions(
    env_with_agent: Any,
) -> None:
    client, srv, auth, aid = env_with_agent
    intruder = await create_user(client, auth, username="workspace-import-intruder")
    srv.services.agent_repo.set_shared(aid, True)
    foreign = await client.post(
        f"/api/agents/{aid}/upload",
        files={"file": ("private-file.txt", b"immutable revision", "text/plain")},
        headers=intruder,
    )
    assert foreign.status_code == 200, foreign.text
    foreign_revision = srv.services.chat_file_repo.get_revision(foreign.json()["revision_id"])
    assert foreign_revision is not None

    workspace = srv.app_runtime.agent_registry.workspace_for_agent(aid)
    assert workspace is not None
    for mode in ("merge", "replace"):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as bundle:
            bundle.writestr(foreign_revision.workspace_path, b"attacker overwrite")
            bundle.writestr(f"imported-{mode}.txt", mode.encode())
        response = await client.post(
            f"/api/agents/{aid}/workspace/archive?mode={mode}",
            files={"file": ("workspace.zip", buf.getvalue(), "application/zip")},
            headers=auth,
        )
        assert response.status_code == 200, response.text
        assert any(
            "preserved 1 registered ChatFile" in warning for warning in response.json()["warnings"]
        )
        assert (
            await workspace.adownload_bytes(foreign_revision.workspace_path)
            == b"immutable revision"
        )
        assert await workspace.adownload_bytes(f"imported-{mode}.txt") == mode.encode()


async def test_delete_draft_only_deletes_owned_unbound_files(env_with_agent: Any) -> None:
    client, srv, auth, aid = env_with_agent
    upload = await client.post(
        f"/api/agents/{aid}/upload",
        files={"file": ("draft.txt", b"draft body", "text/plain")},
        headers=auth,
    )
    assert upload.status_code == 200, upload.text
    draft_id = upload.json()["file_id"]
    url = f"/api/agents/{aid}/files/{draft_id}/draft"

    deleted = await client.delete(url, headers=auth)
    assert deleted.status_code == 204
    row = srv.services.chat_file_repo.get(draft_id)
    assert row is not None and row.deleted_at is not None
    assert (await client.delete(url, headers=auth)).status_code == 404
    assert (
        await client.get(f"/api/agents/{aid}/files/{draft_id}", headers=auth)
    ).status_code == 404

    bound = await _upload_and_bind(client, srv, auth, aid, filename="bound.txt")
    bound_delete = await client.delete(
        f"/api/agents/{aid}/files/{bound['file_id']}/draft", headers=auth
    )
    assert bound_delete.status_code == 404
    bound_row = srv.services.chat_file_repo.get(bound["file_id"])
    assert bound_row is not None and bound_row.deleted_at is None

    foreign_upload = await client.post(
        f"/api/agents/{aid}/upload",
        files={"file": ("foreign.txt", b"owned by admin", "text/plain")},
        headers=auth,
    )
    intruder = await create_user(client, auth, username="draft-delete-intruder")
    foreign_url = f"/api/agents/{aid}/files/{foreign_upload.json()['file_id']}/draft"
    foreign_delete = await client.delete(foreign_url, headers=intruder)
    # The agent-level guard may reject this user before file ownership lookup.
    assert foreign_delete.status_code in (403, 404)
    foreign_row = srv.services.chat_file_repo.get(foreign_upload.json()["file_id"])
    assert foreign_row is not None and foreign_row.deleted_at is None


async def test_read_surface_hides_other_users_files(env_with_agent: Any) -> None:
    client, srv, auth, aid = env_with_agent
    body = await _upload_and_bind(client, srv, auth, aid)
    intruder = await create_user(client, auth, username="intruder")

    detail = await client.get(f"/api/agents/{aid}/files/{body['file_id']}", headers=intruder)
    assert detail.status_code in (403, 404)

    segments = await client.get(
        f"/api/agents/{aid}/files/{body['file_id']}/segments", headers=intruder
    )
    assert segments.status_code in (403, 404)

    search = await client.post(
        f"/api/agents/{aid}/files/{body['file_id']}/search",
        headers=intruder,
        json={"query": "beta"},
    )
    assert search.status_code in (403, 404)


async def test_read_surface_requires_login(env_with_agent: Any) -> None:
    client, _srv, _auth, aid = env_with_agent
    response = await client.get(f"/api/agents/{aid}/files/nope")
    assert response.status_code == 401
